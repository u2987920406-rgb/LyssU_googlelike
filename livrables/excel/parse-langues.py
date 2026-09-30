from openpyxl import load_workbook
wb = load_workbook("/home/raf/projets/ulysse/livrables/excel/langues-les-plus-parlees.xlsx")
print("feuilles:", wb.sheetnames)
ws = wb["Top 10 langues"]
nform = sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
print("formules feuille1:", nform)
print("E5:", ws["E5"].value, "| F5:", ws["F5"].value, "| E15:", ws["E15"].value)
ws2 = wb["Analyse"]
f2 = [c.value for r in ws2.iter_rows() for c in r if isinstance(c.value, str) and c.value.startswith("=")]
print("formules feuille2:", len(f2))
for w in "ABCDEFGHI":
    print("larg", w, ws.column_dimensions[w].width)
print("freeze:", ws.freeze_panes, "| en-tete gras:", ws["B4"].font.bold, "| remplissage:", ws["A4"].fill.fgColor.rgb)
print("L15 label:", ws["B15"].value, "| C15:", ws["C15"].value)
