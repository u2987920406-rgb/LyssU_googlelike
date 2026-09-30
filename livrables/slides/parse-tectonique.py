from pptx import Presentation
p = Presentation("/home/raf/projets/ulysse/livrables/slides/la-tectonique-des-plaques.pptx")
print("slides:", len(p.slides))
for i, s in enumerate(p.slides, 1):
    texts = [sh.text_frame.text.replace("\n", " | ")[:110] for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    print(i, "||", " ~~ ".join(texts)[:380])
