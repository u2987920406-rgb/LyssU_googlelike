#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Parse la-photosynthese.pptx : contrôle structure, débords, remplissage."""
from pptx import Presentation
from pptx.util import Emu

p = Presentation("la-photosynthese.pptx")
print("diapos:", len(p.slides))
W = p.slide_width
H = p.slide_height
for i, s in enumerate(p.slides, 1):
    over = []
    texts = 0
    chars = 0
    for sh in s.shapes:
        if sh.left is not None and sh.width is not None and sh.left + sh.width > W + Emu(10000):
            over.append(("W", sh.shape_id))
        if sh.top is not None and sh.height is not None and sh.top + sh.height > H + Emu(10000):
            over.append(("H", sh.shape_id))
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t:
                texts += 1
                chars += len(t)
    print(f"d{i}: shapes={len(s.shapes)} textboxes={texts} chars={chars} debords={over}")
