#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Génère slides/la-photosynthese.pptx (8 diapos, thème nature/sciences du vivant)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

GREEN = RGBColor(0x0E, 0x2A, 0x1B)   # fond profond
GREEN2 = RGBColor(0x14, 0x38, 0x25)  # bandeau
LEAF = RGBColor(0x5C, 0xD6, 0x7B)    # accent vert clair
GOLD = RGBColor(0xFF, 0xC1, 0x4D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xDD, 0xF2, 0xE2)
GREY = RGBColor(0x9C, 0xC2, 0xA8)
CARD = RGBColor(0x18, 0x42, 0x2C)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W = 13.333


def bg(slide, color=GREEN):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def rect(slide, x, y, w, h, color, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    return sh


def txt(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    first = True
    for text, size, bold, color in runs:
        for line in text.split("\n"):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = align
            p.line_spacing = spacing
            r = p.add_run()
            r.text = line
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = "Calibri"
    return tb


def header(slide, kicker, title):
    rect(slide, 0, 0, W, 1.32, GREEN2)
    rect(slide, 0, 1.32, W, 0.05, LEAF)
    txt(slide, 0.55, 0.14, 12.2, 0.4, [(kicker, 13, True, LEAF)])
    txt(slide, 0.55, 0.46, 12.2, 0.8, [(title, 30, True, WHITE)])


def bullets(slide, x, y, w, h, items, size=16, color=LIGHT, gap=8):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
        r = p.add_run()
        r.text = "▪  " + it
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = "Calibri"
    return tb


def src(slide, text):
    txt(slide, 0.7, 6.85, 11.9, 0.4, [(text, 11, False, GREY)])


# ───────────────────────── 1. Titre ─────────────────────────
s = prs.slides.add_slide(BLANK)
bg(s)
rect(s, 0, 2.9, W, 0.06, LEAF)
txt(s, 1.0, 1.75, 11.3, 0.5, [("SCIENCES DU VIVANT", 16, True, LEAF)], align=PP_ALIGN.CENTER)
txt(s, 1.0, 2.15, 11.3, 0.75, [("La photosynthèse", 42, True, WHITE)], align=PP_ALIGN.CENTER)
txt(s, 1.0, 3.15, 11.3, 1.1,
    [("Comment les plantes transforment la lumière en matière vivante", 20, False, LIGHT),
     ("et fabriquent le dioxygène que nous respirons", 20, False, LIGHT)],
    align=PP_ALIGN.CENTER, spacing=1.1)
txt(s, 1.0, 6.45, 11.3, 0.5,
    [("Sources vérifiées citées dans chaque diapo · Wikipédia · ENS Lyon · PMC · NOAA · Britannica · NASA", 12, False, GREY)],
    align=PP_ALIGN.CENTER)

# ───────────────────────── 2. Définition + équation ─────────────────────────
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "DÉFINITION", "Qu'est-ce que la photosynthèse ?")
bullets(s, 0.7, 1.8, 11.9, 2.7, [
    "Du grec « photo » (lumière) et « synthèse » (assemblage) : fabriquer de la matière organique avec la lumière.",
    "Les organismes autotrophes — plantes, algues, cyanobactéries — captent l'énergie lumineuse et la convertissent en énergie chimique (glucides).",
    "C'est la base de presque toute la chaîne alimentaire terrestre et la source du dioxygène atmosphérique.",
], size=18, gap=12)
rect(s, 1.4, 4.75, 10.5, 1.35, CARD, line=LEAF)
txt(s, 1.6, 4.95, 10.1, 1.0,
    [("Équation bilan simplifiée :", 15, True, LEAF),
     ("6 CO₂ + 6 H₂O  +  lumière  →  C₆H₁₂O₆ (glucose)  +  6 O₂", 22, True, GOLD)],
    align=PP_ALIGN.CENTER, spacing=1.15)
src(s, "Sources : Wikipédia « Photosynthèse » (définition, équation) ; Planet-Terre ENS- Lyon (bilan de masse : 44 g de CO₂ → 12 g de C fixé + 32 g d'O₂) ; Encyclopédie de l'environnement « La biosphère ».")

# ───────────────────────── 3. Deux étapes ─────────────────────────
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "MÉCANISME", "Un processus en deux temps")
rect(s, 0.7, 1.85, 5.85, 4.35, CARD, line=LEAF)
txt(s, 0.95, 2.05, 5.35, 0.55, [("1. Phase photochimique (« claire »)", 19, True, LEAF)])
bullets(s, 0.95, 2.7, 5.35, 3.3, [
    "Dans la membrane des thylakoïdes (chloroplaste).",
    "Les photons excitent la chlorophylle ; l'énergie sert à la photolyse de l'eau : H₂O → O₂ + électrons + protons.",
    "Résultat : deux « batteries » chimiques, l'ATP et le NADPH — et le dioxygène comme sous-produit.",
], size=15, gap=9)
rect(s, 6.8, 1.85, 5.85, 4.35, CARD, line=GOLD)
txt(s, 7.05, 2.05, 5.35, 0.55, [("2. Cycle de Calvin (« sombre »)", 19, True, GOLD)])
bullets(s, 7.05, 2.7, 5.35, 3.3, [
    "Dans le stroma du chloroplaste — indépendant de la lumière, mais il en consomme le fruit.",
    "L'ATP et le NADPH dépensés permettent de fixer le CO₂ sur du RuBP et de bâtir du glucose.",
    "Le glucose est ensuite polymérisé en amidon (réserve) et cellulose (parois des cellules).",
], size=15, gap=9)
src(s, "Sources : Wikipédia « Photosynthèse » (phase photochimique / cycle de Calvin) ; PMC « Photosynthesis » (PMC5264509) ; S. Merchant et al., revue PMC7489092 (Calvin-Benson, ATP/NADPH).")

# ───────────────────────── 4. Le chloroplaste ─────────────────────────
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "AU CŒUR DE LA CELLULE", "Le chloroplaste et la chlorophylle")
bullets(s, 0.7, 1.85, 11.9, 4.6, [
    "Les feuilles vertes abritent des chloroplastes : organites à double membrane, capteurs de lumière de la cellule.",
    "Leurs pigments — chlorophylles et caroténoïdes — absorbent surtout le bleu et le rouge, et réfléchissent le vert : d'où la couleur des feuilles.",
    "Une unité photosynthétique implique environ 600 molécules de chlorophylle (travaux d'Emerson et Arnold, années 1930-1940).",
    "Au moins 8 photons sont nécessaires pour réduire une molécule de CO₂ — l'exigence quantique de la photosynthèse.",
    "L'essentiel de la masse sèche d'un arbre ne vient pas du sol, mais du CO₂ de l'air fixé par la photosynthèse.",
], size=18, gap=12)
src(s, "Sources : longuevergne.free.fr « La photosynthèse et la chlorophylle » (8 photons/CO₂, unité photosynthétique 600 chlorophylles, Emerson & Arnold) ; Wikipédia « Photosynthèse » (pigments, chloroplastes) ; Planet-Terre ENS-Lyon (masse sèche et CO₂).")

# ───────────────────────── 5. Chiffres clés ─────────────────────────
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "EN UN COUP D'ŒIL", "La photosynthèse en chiffres")
cards = [
    ("≈ 200 Gt", "de CO₂ converties en composés organiques chaque année"),
    ("≈ 140 Gt", "de dioxygène rejetées dans l'atmosphère par an"),
    ("≈ 50 %", "de l'O₂ terrestre produit par l'océan (phytoplancton)"),
    ("jusqu'à 20 %", "de l'O₂ de la biosphère dû au seul Prochlorococcus"),
    ("8", "photons minimum par molécule de CO₂ réduite"),
    ("≈ 30 %", "rendement maximal de la conversion lumineuse"),
]
x0, y0, cw, ch, gx, gy = 0.7, 1.8, 3.95, 2.15, 0.18, 0.35
for i, (big, small) in enumerate(cards):
    cx = x0 + (i % 3) * (cw + gx)
    cy = y0 + (i // 3) * (ch + gy)
    rect(s, cx, cy, cw, ch, CARD, line=LEAF)
    txt(s, cx + 0.1, cy + 0.28, cw - 0.2, 0.85, [(big, 26, True, GOLD)], align=PP_ALIGN.CENTER)
    txt(s, cx + 0.15, cy + 1.15, cw - 0.3, 0.95, [(small, 14, False, LIGHT)], align=PP_ALIGN.CENTER)
src(s, "Sources : PMC « Photosynthesis » (PMC5264509 : ≈200 Gt CO₂/an, ≈140 Gt O₂/an) ; NOAA Ocean Service (≈50 % de l'O₂ vient de l'océan, Prochlorococcus jusqu'à 20 %) ; longuevergne.free.fr (8 photons, rendement ~30 %).")

# ───────────────────────── 6. Histoire des découvertes ─────────────────────────
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "REPÈRES", "Trois siècles de découvertes")
rect(s, 1.15, 2.6, 0.06, 3.9, LEAF)
events = [
    ("1771-1777", "Joseph Priestley montre qu'une plante « répare » l'air vicié par une bougie (rejet d'oxygène)."),
    ("1779", "Jan Ingenhousz établit le rôle essentiel de la lumière et des parties vertes de la plante."),
    ("1782", "Jean Senebier montre que les plantes consomment le dioxyde de carbone pour croître."),
    ("1804", "Nicolas de Saussure prouve que l'eau est un réactif : la plante gagne en masse par absorption de CO₂ et d'eau."),
    ("1862", "Julius von Sachs montre la synthèse de l'amidon dans les feuilles éclairées."),
    ("1961", "Melvin Calvin reçoit le prix Nobel de chimie pour le cycle de fixation du carbone (cycle de Calvin)."),
]
y = 1.62
for date, desc in events:
    rect(s, 1.02, y + 0.12, 0.3, 0.3, GOLD)
    txt(s, 1.55, y, 1.65, 0.5, [(date, 16, True, LEAF)])
    txt(s, 3.3, y, 9.15, 0.85, [(desc, 15, False, LIGHT)], spacing=1.02)
    y += 0.86
src(s, "Sources : Britannica « History of the study of photosynthesis » (Priestley 1771, Ingenhousz 1779, de Saussure 1804) ; Wikipédia « Photosynthèse » (Senebier, Sachs) ; PMC7489092 (Calvin, Nobel 1961, Benson-Bassham).")

# ───────────────────────── 7. Impact planétaire ─────────────────────────
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "ÉCHELLE PLANÉTAIRE", "L'oxygène d'une planète : l'événement de la Grande Oxydation")
bullets(s, 0.7, 1.85, 11.9, 4.6, [
    "Il y a ≈ 2,7 milliards d'années, les cyanobactéries inventent la photosynthèse oxygénique : l'eau devient « carburant », l'O₂ est un déchet.",
    "Grande Oxydation (≈ 2,4 milliards d'années) : l'oxygène s'accumule dans l'atmosphère et atteint ≈ 1 % du niveau actuel en 600 millions d'années.",
    "Cet événement a été une extinction massive pour la vie anaérobie, mais a permis le métabolisme aérobie et la vie multicellulaire.",
    "Aujourd'hui, la pompe biologique et l'enfouissement géologique du carbone maintiennent l'équilibre O₂/CO₂ hérité de ce processus.",
    "Le phytoplancton océanique assure à lui seul presque la moitié de la production primaire mondiale.",
], size=17, gap=11)
src(s, "Sources : NASA Astrobiology « Revisiting Earth's Oxygenation 2.4 Billion Years Ago » ; Britannica « Evolution of the atmosphere » (GOE ≈ 2,4 Ga, ≈ 1 % des niveaux actuels sur 600 Ma) ; ASM « The Great Oxidation Event » (cyanobactéries ≈ 2,7 Ga) ; NASA My NASA Data (phytoplancton ≈ 50 % de la production primaire).")

# ───────────────────────── 8. Sources ─────────────────────────
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "RÉFÉRENCES", "Sources consultées")
bullets(s, 0.7, 1.8, 12.0, 5.0, [
    "Wikipédia — « Photosynthèse » (définition, équation-bilan, phase photochimique / cycle de Calvin, pigments ; Hales 1727, Priestley, Ingenhousz 1779, Senebier, Sachs 1862).",
    "Planet-Terre ENS-Lyon — « L'Amazonie, le poumon de la Terre… » (bilan de masse : 44 g CO₂ → 12 g C + 32 g O₂).",
    "PMC — « Photosynthesis » (PMC5264509) : ≈200 Gt de CO₂ et ≈140 Gt de O₂ par an.",
    "PMC7489092 — « Photosynthesis: basics, history and modelling » (Priestley, Ingen-Housz, Senebier, Sachs, Nobel de Melvin Calvin 1961).",
    "NOAA Ocean Service — « How much oxygen comes from the ocean? » (≈50 % de l'O₂ terrestre, Prochlorococcus jusqu'à 20 %).",
    "Britannica — « History of the study of photosynthesis » et « Evolution of the atmosphere » (Grande Oxydation ≈ 2,4 Ga).",
    "NASA Astrobiology (2,4 Ga) ; ASM « The Great Oxidation Event » (cyanobactéries ≈ 2,7 Ga) ; NASA My NASA Data (phytoplancton).",
], size=15, gap=9)

prs.save("/home/raf/projets/ulysse/livrables/slides/la-photosynthese.pptx")
print("OK 8 diapos")
